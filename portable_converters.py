"""Conversion engines installed with pip; no Office, LibreOffice or Poppler needed."""
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape


def pdf_to_jpegs(source, output):
    """Render one page at a time to keep memory bounded; PDFium ships in wheels."""
    import pypdfium2 as pdfium
    images = []
    with pdfium.PdfDocument(str(source)) as document:
        for index in range(len(document)):
            page = document[index]
            try:
                longest = max(page.get_size())
                if longest <= 0:
                    raise ValueError('This PDF contains an invalid page size.')
                bitmap = page.render(scale=min(150 / 72, 2400 / longest))
                try:
                    image = bitmap.to_pil()
                    try:
                        path = output / f'page-{index + 1:03}.jpg'
                        image.convert('RGB').save(path, quality=90)
                        images.append(path)
                    finally:
                        image.close()
                finally:
                    bitmap.close()
            finally:
                page.close()
    return images


def word_to_pdf(source, destination):
    """Reflow DOCX text, basic styles, tables and embedded raster images into PDF.

    This is intentionally a portable renderer, not an exact Word layout engine.
    Floating objects, headers, footers and complex pagination are not reproduced.
    """
    import reportlab
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph as WordParagraph
    from docx.text.run import Run
    from docx.table import Table as WordTable
    from PIL import Image as PILImage
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT, TA_JUSTIFY
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Image, LongTable, TableStyle

    # Vera is distributed with ReportLab, so ordinary Latin text works without
    # downloading fonts or relying on a particular operating system's fonts.
    font_dir = Path(reportlab.__file__).parent / 'fonts'
    faces = {'Portable': 'Vera.ttf', 'Portable-Bold': 'VeraBd.ttf',
             'Portable-Italic': 'VeraIt.ttf', 'Portable-BoldItalic': 'VeraBI.ttf'}
    for name, filename in faces.items():
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(font_dir / filename)))
    pdfmetrics.registerFontFamily('Portable', normal='Portable', bold='Portable-Bold',
                                  italic='Portable-Italic', boldItalic='Portable-BoldItalic')
    supported = pdfmetrics.getFont('Portable').face.charToGlyph
    missing_glyphs = False
    skipped_images = False
    document = Document(source)
    section = document.sections[0]
    page_width = float(section.page_width.pt) if section.page_width else 595.3
    page_height = float(section.page_height.pt) if section.page_height else 841.9
    # Reject pathological geometry instead of allocating huge pages or frames.
    if not (144 <= page_width <= 2000 and 144 <= page_height <= 2000):
        raise ValueError('This document has an unsupported page size.')
    left = min(float(section.left_margin.pt), page_width / 4) if section.left_margin else 54
    right = min(float(section.right_margin.pt), page_width / 4) if section.right_margin else 54
    top = min(float(section.top_margin.pt), page_height / 4) if section.top_margin else 54
    bottom = min(float(section.bottom_margin.pt), page_height / 4) if section.bottom_margin else 54
    width = page_width - left - right - 12
    height = page_height - top - bottom - 24
    base = ParagraphStyle('Body', fontName='Portable', fontSize=10, leading=15,
                          spaceAfter=8, textColor=colors.HexColor('#17252b'))
    alignment = {0: TA_LEFT, 1: TA_CENTER, 2: TA_RIGHT, 3: TA_JUSTIFY}

    def safe_text(text):
        nonlocal missing_glyphs
        if any(ord(c) not in supported for c in text if not c.isspace()):
            missing_glyphs = True
        return escape(text).replace('\n', '<br/>').replace('\t', '    ')

    def paragraph_flows(paragraph, available_width, in_table=False):
        nonlocal skipped_images
        style_name = paragraph.style.name if paragraph.style else ''
        size = 10
        if style_name.startswith('Heading'):
            size = {'Heading 1': 20, 'Heading 2': 16, 'Heading 3': 13}.get(style_name, 12)
        elif style_name in ('Title', 'Subtitle'):
            size = 24 if style_name == 'Title' else 14
        style = ParagraphStyle('Content', parent=base, fontSize=size, leading=size * 1.45,
                               alignment=alignment.get(paragraph.alignment, TA_LEFT),
                               fontName='Portable-Bold' if style_name.startswith('Heading') else 'Portable',
                               keepWithNext=style_name.startswith('Heading') and not in_table)
        flows, fragments = [], []
        numbered = 'List' in style_name or bool(paragraph._p.xpath('./w:pPr/w:numPr'))

        def flush():
            if fragments:
                markup = ''.join(fragments)
                flows.append(Paragraph(markup, style, bulletText='•' if numbered else None))
                fragments.clear()

        if paragraph.paragraph_format.page_break_before and not in_table:
            flows.append(PageBreak())
        # Descendant runs include hyperlink text, in document order.
        for element in paragraph._p.xpath('.//w:r'):
            run = Run(element, paragraph)
            text = safe_text(run.text)
            if text:
                if run.bold:
                    text = '<b>' + text + '</b>'
                if run.italic:
                    text = '<i>' + text + '</i>'
                if run.underline:
                    text = '<u>' + text + '</u>'
                if run.font.size:
                    text = f'<font size="{max(6, min(48, run.font.size.pt))}">{text}</font>'
                fragments.append(text)
            for blip in element.xpath('.//a:blip'):
                relationship = blip.get(qn('r:embed'))
                if not relationship:
                    skipped_images = True
                    continue
                flush()
                try:
                    # Only embedded bytes are read; never fetch external image URLs.
                    part = document.part.related_parts[relationship]
                    with PILImage.open(BytesIO(part.blob)) as source_image:
                        source_image.load()
                        stream = BytesIO()
                        source_image.convert('RGBA').save(stream, 'PNG')
                        iw, ih = source_image.size
                    stream.seek(0)
                    extents = element.xpath('.//wp:extent')
                    image_width = float(extents[0].get('cx')) / 12700 if extents else iw * .75
                    image_height = float(extents[0].get('cy')) / 12700 if extents else ih * .75
                    if image_width <= 0 or image_height <= 0:
                        raise ValueError('Invalid image dimensions')
                    ratio = min(1, available_width / image_width, height / image_height)
                    flows.append(Image(stream, width=image_width * ratio, height=image_height * ratio))
                    flows.append(Spacer(1, 8))
                except Exception:
                    skipped_images = True
                    flows.append(Paragraph('[Image could not be rendered]', base))
            if element.xpath('./w:br[@w:type="page"]') and not in_table:
                flush()
                flows.append(PageBreak())
        flush()
        return flows or [Spacer(1, 6)]

    story = []
    for element in document.element.body:
        if element.tag == qn('w:p'):
            story.extend(paragraph_flows(WordParagraph(element, document), width))
        elif element.tag == qn('w:tbl'):
            table = WordTable(element, document)
            if not table.rows or not table.columns:
                continue
            column_width = width / len(table.columns)
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    flows = []
                    for paragraph in cell.paragraphs:
                        flows.extend(paragraph_flows(paragraph, max(12, column_width - 12), True))
                    cells.append(flows or [Paragraph('', base)])
                rows.append(cells)
            rendered = LongTable(rows, colWidths=[column_width] * len(table.columns),
                                 hAlign='LEFT', splitByRow=1, splitInRow=1)
            rendered.setStyle(TableStyle([
                ('GRID', (0, 0), (-1, -1), .5, colors.HexColor('#ced5dd')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 6),
                ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
            story.extend([rendered, Spacer(1, 12)])
    if not story:
        story = [Paragraph(' ', base)]
    SimpleDocTemplate(str(destination), pagesize=(page_width, page_height),
        leftMargin=left, rightMargin=right, topMargin=top, bottomMargin=bottom,
        title='Converted document').build(story)
    message = ('Portable Word conversion: text, basic styling, tables and embedded images are reflowed. '
               'Fonts, page layout, numbering, headers, footers and complex objects may differ or be omitted.')
    if skipped_images:
        message += ' Some images could not be rendered.'
    if missing_glyphs:
        message += ' Some characters are not supported by the bundled font; check the PDF before sharing.'
    return message
