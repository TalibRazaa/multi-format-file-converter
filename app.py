"""Small Flask converter. Run with `python app.py` or the supplied Dockerfile."""
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import threading
import time
import uuid
import zipfile

from flask import Flask, jsonify, render_template, request, send_file
from werkzeug.exceptions import HTTPException, RequestEntityTooLarge

BASE = Path(__file__).resolve().parent
UPLOADS = BASE / 'uploads'
CONVERTED = BASE / 'converted'
MAX_BYTES = 10 * 1024 * 1024
TTL = 3600
TIMEOUT = 120
MODES = {'pdf': ['pdf-jpeg', 'pdf-docx'], 'jpg': ['image-pdf'],
         'jpeg': ['image-pdf'], 'png': ['image-pdf'], 'docx': ['docx-pdf']}
for folder in (UPLOADS, CONVERTED):
    folder.mkdir(exist_ok=True)

app = Flask(__name__)
# Permit multipart overhead, then independently enforce the exact file limit.
app.config.update(MAX_CONTENT_LENGTH=MAX_BYTES + 128 * 1024,
                  MAX_FORM_MEMORY_SIZE=256 * 1024, MAX_FORM_PARTS=8)
slots = threading.BoundedSemaphore(2)


def job_path(token):
    if not isinstance(token, str) or not re.fullmatch(r'[a-f0-9]{32}', token):
        raise ValueError('Invalid file reference. Upload the file again.')
    folder = UPLOADS / token
    if not folder.is_dir() or time.time() - folder.stat().st_mtime > TTL:
        raise ValueError('This file has expired. Please upload it again.')
    return folder


def cleanup():
    """Remove entire jobs after an hour, including uploads never converted."""
    for root in (UPLOADS, CONVERTED):
        for folder in root.iterdir():
            try:
                if folder.is_dir() and time.time() - folder.stat().st_mtime > TTL:
                    shutil.rmtree(folder, ignore_errors=True)
            except FileNotFoundError:
                pass


def janitor():
    while True:
        cleanup()
        time.sleep(60)


@app.after_request
def secure_response(response):
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['Cache-Control'] = 'no-store'
    response.headers['Content-Security-Policy'] = (
        "default-src 'self'; script-src 'self'; style-src 'self'; "
        "img-src 'self' data:; object-src 'none'; base-uri 'self'; frame-ancestors 'none'")
    return response


@app.errorhandler(RequestEntityTooLarge)
def too_large(_):
    return jsonify(error='The file is too large. Maximum size is 10 MB.'), 413


@app.errorhandler(ValueError)
def invalid(error):
    return jsonify(error=str(error)), 400


@app.errorhandler(HTTPException)
def http_error(error):
    return jsonify(error=error.description), error.code


@app.errorhandler(Exception)
def unexpected(error):
    app.logger.exception('Unexpected request failure')
    return jsonify(error='Something went wrong. Please try again.'), 500


@app.get('/')
def index():
    return render_template('index.html')


# --- TEMPORARY DEBUG ROUTE: remove this once the template issue is solved ---
@app.get('/debug-files')
def debug_files():
    return jsonify(
        base_dir=str(BASE),
        base_contents=os.listdir(BASE),
        templates_exists=(BASE / 'templates').exists(),
        templates_contents=os.listdir(BASE / 'templates') if (BASE / 'templates').exists() else 'MISSING'
    )
# --- END TEMPORARY DEBUG ROUTE ---


@app.post('/upload')
def upload():
    # Never use a supplied filename as a filesystem path.
    incoming = request.files.get('file')
    if not incoming or not incoming.filename:
        raise ValueError('Choose a file first.')
    ext = incoming.filename.rsplit('.', 1)[-1].lower()
    if ext not in MODES:
        raise ValueError('Supported files: PDF, JPEG, PNG and DOCX.')
    token = uuid.uuid4().hex
    folder = UPLOADS / token
    folder.mkdir(mode=0o700)
    try:
        total = 0
        with (folder / ('input.' + ext)).open('wb') as target:
            while chunk := incoming.stream.read(64 * 1024):
                total += len(chunk)
                if total > MAX_BYTES:
                    raise RequestEntityTooLarge()
                target.write(chunk)
        if total == 0:
            raise ValueError('The file is empty.')
        (folder / 'meta.json').write_text(json.dumps({'ext': ext}))
        return jsonify(id=token, conversions=MODES[ext]), 201
    except Exception:
        shutil.rmtree(folder, ignore_errors=True)
        raise


@app.post('/convert')
def convert():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        raise ValueError('Expected a file reference and conversion type.')
    folder = job_path(data.get('id'))
    mode = data.get('type')
    ext = json.loads((folder / 'meta.json').read_text())['ext']
    if mode not in MODES[ext]:
        raise ValueError('That conversion does not match this file.')
    if not slots.acquire(blocking=False):
        return jsonify(error='The converter is busy. Please try again shortly.'), 503
    locked = False
    output = CONVERTED / folder.name
    try:
        # Exclusive creation also prevents two server processes using the same job.
        try:
            (folder / 'busy').touch(exist_ok=False)
            locked = True
        except FileExistsError:
            return jsonify(error='This file is already converting.'), 409
        if (output / 'result.json').exists():
            raise ValueError('This file has already been converted. Upload it again for another format.')
        os.utime(folder, None)
        output.mkdir(exist_ok=True, mode=0o700)
        # Isolate parsing/conversion in a bounded child process. User input never
        # becomes a shell command; subprocess runs with shell=False by default.
        process = subprocess.Popen(
            [sys.executable, str(BASE / 'app.py'), '--worker', folder.name, mode],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            start_new_session=(os.name == 'posix'))
        try:
            code = process.wait(timeout=TIMEOUT)
        except subprocess.TimeoutExpired:
            if os.name == 'posix':
                import signal
                os.killpg(process.pid, signal.SIGKILL)
            else:
                process.kill()
            process.wait()
            shutil.rmtree(output, ignore_errors=True)
            return jsonify(error='Conversion took too long. Try a smaller document.'), 422
        result_file = output / 'result.json'
        if code or not result_file.exists():
            error_file = output / 'error.txt'
            message = error_file.read_text(encoding='utf-8') if error_file.exists() else 'Unable to convert this file.'
            shutil.rmtree(output, ignore_errors=True)
            return jsonify(error=message), 422
        result = json.loads(result_file.read_text())
        # Delete the source immediately after success; output is retained for one hour.
        (folder / ('input.' + ext)).unlink(missing_ok=True)
        return jsonify(download='/download/' + folder.name,
                       filename=result['filename'], warning=result.get('warning', ''))
    finally:
        if locked:
            (folder / 'busy').unlink(missing_ok=True)
            if folder.exists():
                os.utime(folder, None)
        slots.release()


@app.get('/download/<token>')
def download(token):
    # Opaque 128-bit job IDs act as private, short-lived download capabilities.
    job_path(token)
    folder = CONVERTED / token
    manifest = folder / 'result.json'
    if not manifest.exists() or time.time() - manifest.stat().st_mtime > TTL:
        return jsonify(error='Download expired or is not ready. Upload the file again.'), 404
    result = json.loads(manifest.read_text())
    # Preserve the file until TTL so streaming, interrupted transfers and retries work.
    return send_file(folder / result['filename'], as_attachment=True,
                     download_name=result['filename'], max_age=0)


def run_worker(token, mode):
    """Validate actual contents and convert inside the short-lived child."""
    from PIL import Image, ImageOps
    from docx import Document
    from pypdf import PdfReader
    Image.MAX_IMAGE_PIXELS = 20_000_000
    import warnings
    warnings.simplefilter('error', Image.DecompressionBombWarning)
    source_dir = job_path(token)
    out = CONVERTED / token
    ext = json.loads((source_dir / 'meta.json').read_text())['ext']
    source = source_dir / ('input.' + ext)
    warning = ''
    try:
        if ext == 'pdf':
            if not source.read_bytes()[:1024].lstrip().startswith(b'%PDF-'):
                raise ValueError('This file is not a valid PDF.')
            reader = PdfReader(source)
            if reader.is_encrypted:
                raise ValueError('Password-protected PDFs are not supported. Upload an unlocked copy.')
            if not 1 <= len(reader.pages) <= 100:
                raise ValueError('Use a PDF with between 1 and 100 pages.')
        if mode == 'image-pdf':
            with Image.open(source) as original:
                if original.format not in ('JPEG', 'PNG'):
                    raise ValueError('The image contents must be JPEG or PNG.')
                original.load()
                oriented = ImageOps.exif_transpose(original).convert('RGBA')
                # Flatten transparency onto white and respect camera orientation.
                flat = Image.new('RGB', oriented.size, 'white')
                flat.paste(oriented, mask=oriented.getchannel('A'))
                flat.save(out / 'converted.pdf', 'PDF', resolution=150)
            filename = 'converted.pdf'
        elif mode == 'pdf-jpeg':
            from portable_converters import pdf_to_jpegs
            images = pdf_to_jpegs(source, out)
            if len(images) == 1:
                Path(images[0]).rename(out / 'converted.jpg')
                filename = 'converted.jpg'
            else:
                filename = 'converted-pages.zip'
                with zipfile.ZipFile(out / filename, 'w', zipfile.ZIP_DEFLATED) as archive:
                    for number, path in enumerate(images, 1):
                        archive.write(path, f'page-{number:03}.jpg')
                        Path(path).unlink()
        elif mode == 'pdf-docx':
            document = Document()
            has_text = False
            blank_pages = 0
            for index, page in enumerate(reader.pages):
                text = page.extract_text() or ''
                if index:
                    document.add_page_break()
                if text.strip():
                    has_text = True
                    for line in text.splitlines():
                        document.add_paragraph(line)
                else:
                    blank_pages += 1
                    document.add_paragraph('[No extractable text on this page; OCR may be required.]')
            if not has_text:
                raise ValueError('This PDF has no extractable text. Scanned PDFs need OCR, which is not included.')
            warning = 'Editable text only. Images, tables and original layout are not preserved.'
            if blank_pages:
                warning += f' {blank_pages} page(s) had no extractable text and were marked in the output.'
            filename = 'converted.docx'
            document.save(out / filename)
        elif mode == 'docx-pdf':
            # Check the OOXML package before opening it in the portable renderer.
            with zipfile.ZipFile(source) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(x.file_size for x in members) > 50 * 1024 * 1024:
                    raise ValueError('This Word file expands beyond the supported size limit.')
                names = archive.namelist()
                if 'word/document.xml' not in names or '[Content_Types].xml' not in names:
                    raise ValueError('This file is not a valid DOCX document.')
                if any('vbaProject' in name for name in names):
                    raise ValueError('Macro-enabled documents are not supported.')
            from portable_converters import word_to_pdf
            filename = 'converted.pdf'
            warning = word_to_pdf(source, out / filename)
        else:
            raise ValueError('Unsupported conversion.')
        (out / 'result.json').write_text(json.dumps({'filename': filename, 'warning': warning}))
    except ValueError as error:
        (out / 'error.txt').write_text(str(error), encoding='utf-8')
        return 1
    except Exception:
        (out / 'error.txt').write_text('The file is damaged, unsupported, or could not be processed. Try opening and saving it again.')
        return 1
    return 0


if __name__ == '__main__' and len(sys.argv) == 4 and sys.argv[1] == '--worker':
    raise SystemExit(run_worker(sys.argv[2], sys.argv[3]))

# A periodic sweep works even when no further requests arrive. Each WSGI worker
# may run a sweep safely; expired folders are removed on the next startup too.
cleanup()
threading.Thread(target=janitor, daemon=True).start()

if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5000, debug=False)